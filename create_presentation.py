import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Tree & Earth Color Palette
COLOR_DARK_FOREST  = RGBColor(0x0E, 0x22, 0x13)     # #0E2213 Deep pine/forest
COLOR_BG_PARCHMENT = RGBColor(0xFA, 0xF8, 0xF5)     # #FAF8F5 Soft warm linen
COLOR_SAGE_ACCENT  = RGBColor(0x43, 0x7A, 0x3B)     # #437A3B Vibrant sage
COLOR_TERRACOTTA   = RGBColor(0xBC, 0x6C, 0x25)     # #BC6C25 Warm terracotta/clay
COLOR_CARD_DARK    = RGBColor(0x18, 0x36, 0x1D)     # Forest card
COLOR_CARD_LIGHT   = RGBColor(0xFF, 0xFF, 0xFF)     # Pure white card
COLOR_BORDER_LIGHT = RGBColor(0xE2, 0xDC, 0xD3)     # Subtle warm stone
COLOR_TEXT_DARK    = RGBColor(0x19, 0x23, 0x1D)     # Deep obsidian bark

prs = Presentation()
prs.slide_width = SLIDE_WIDTH
prs.slide_height = SLIDE_HEIGHT
blank_layout = prs.slide_layouts[6]

def add_background(slide, bg_color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background()
    return bg

# ==================== SLIDE 1: THE EXECUTIVE HOOK & PROBLEM ====================
slide1 = prs.slides.add_slide(blank_layout)
add_background(slide1, COLOR_DARK_FOREST)

# Header Tag
tag1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.55), Inches(4.5), Inches(0.38))
tag1.fill.solid()
tag1.fill.fore_color.rgb = RGBColor(0x23, 0x4D, 0x20)
tag1.line.fill.background()
tf_tag1 = tag1.text_frame
tf_tag1.word_wrap = True
p_tag1 = tf_tag1.paragraphs[0]
p_tag1.text = "🌲 URBAN AI INFRASTRUCTURE • HALIFAX REGIONAL MUNICIPALITY"
p_tag1.font.size = Pt(9.5)
p_tag1.font.bold = True
p_tag1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Main Title & Subtitle Box
tb_title1 = slide1.shapes.add_textbox(Inches(0.9), Inches(1.05), Inches(11.5), Inches(1.5))
tf1 = tb_title1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p1_1 = tf1.paragraphs[0]
p1_1.text = "CanopyGuard: Which Tree Falls First?"
p1_1.font.size = Pt(36)
p1_1.font.bold = True
p1_1.font.color.rgb = RGBColor(0xFA, 0xF8, 0xF5)
p1_1.space_after = Pt(4)

p1_2 = tf1.add_paragraph()
p1_2.text = "Predictive Urban Hazard Triage Protecting Cities, Power Grids & Human Lives"
p1_2.font.size = Pt(17)
p1_2.font.color.rgb = RGBColor(0xC2, 0xDE, 0xC0)
p1_2.space_after = Pt(10)

p1_3 = tf1.add_paragraph()
p1_3.text = "Presented by: Devang Jalag • Ketan Khunti • Sanif • Tarun  |  Team Lambda Legends"
p1_3.font.size = Pt(12)
p1_3.font.bold = True
p1_3.font.color.rgb = RGBColor(0xEC, 0xF6, 0xEA)

# 3 Pillars on Slide 1
def make_card(slide, left, top, width, height, pill_text, pill_bg, pill_color, title, items, is_dark=False):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_DARK if is_dark else COLOR_CARD_LIGHT
    card.line.color.rgb = COLOR_SAGE_ACCENT if is_dark else COLOR_BORDER_LIGHT
    card.line.width = Pt(1.5)

    pad = Inches(0.28)
    tb = slide.shapes.add_textbox(left + pad, top + pad, width - (pad * 2), height - (pad * 2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_pill = tf.paragraphs[0]
    p_pill.text = f"[{pill_text.upper()}]"
    p_pill.font.size = Pt(10)
    p_pill.font.bold = True
    p_pill.font.color.rgb = pill_color
    p_pill.space_after = Pt(4)

    p_title = tf.add_paragraph()
    p_title.text = title
    p_title.font.size = Pt(16)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if is_dark else COLOR_DARK_FOREST
    p_title.space_after = Pt(12)

    for item in items:
        p = tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(11.5)
        p.font.color.rgb = RGBColor(0xDF, 0xEE, 0xDC) if is_dark else COLOR_TEXT_DARK
        p.space_after = Pt(8)

make_card(
    slide1, Inches(0.9), Inches(2.9), Inches(3.65), Inches(4.1),
    "$12M+ Municipal Risk", RGBColor(0x4A, 0x28, 0x10), RGBColor(0xF5, 0xB3, 0x82),
    "⚠️ Blind Backlog Crisis",
    [
        "290+ unranked complaints stuck in obsolete first-come FIFO queue",
        "Lethal deadwood over schoolyards waits behind cosmetic branch pruning",
        "Atlantic gale storms turn ignored trunk splits into fatal power blackouts",
        "Cities face millions in wrongful death & infrastructure liability"
    ],
    is_dark=True
)

make_card(
    slide1, Inches(4.84), Inches(2.9), Inches(3.65), Inches(4.1),
    "85% Efficiency Gain", RGBColor(0x1B, 0x45, 0x20), RGBColor(0xA4, 0xDE, 0x9D),
    "⚡ Instant Automated Triage",
    [
        "Replaces 3-week backlog with sub-second priority scoring",
        "Dynamic 0–100 scale: 50% Danger, 25% Wait, 15% Location, 10% Review",
        "Real HRM open data synchronized live across 80,000+ public trees",
        "One-click printable arborist incident posters for immediate truck rollout"
    ],
    is_dark=True
)

make_card(
    slide1, Inches(8.78), Inches(2.9), Inches(3.65), Inches(4.1),
    "Unbreakable Moat", RGBColor(0x4E, 0x2F, 0x12), RGBColor(0xFA, 0xC0, 0x90),
    "👁 Anti-Gaming Vision AI",
    [
        "LlamaParse + Vision LLM extracts arborist defect taxonomy",
        "80% photo weighting stops panic-word queue jumpers",
        "Cross-examines claim text against visual root & wire proof",
        "Unsure guardrails prevent hallucinations before dispatch"
    ],
    is_dark=True
)

# ==================== SLIDE 2: BUSINESS VALUE & SHAREHOLDER ROI ====================
slide2 = prs.slides.add_slide(blank_layout)
add_background(slide2, COLOR_BG_PARCHMENT)

# Header Tag
tag2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.55), Inches(3.2), Inches(0.38))
tag2.fill.solid()
tag2.fill.fore_color.rgb = RGBColor(0xFC, 0xEE, 0xDE)
tag2.line.color.rgb = COLOR_TERRACOTTA
tf_tag2 = tag2.text_frame
tf_tag2.word_wrap = True
p_tag2 = tf_tag2.paragraphs[0]
p_tag2.text = "📈 SCALE & SHAREHOLDER VALUE"
p_tag2.font.size = Pt(9.5)
p_tag2.font.bold = True
p_tag2.font.color.rgb = COLOR_TERRACOTTA

# Title
tb_title2 = slide2.shapes.add_textbox(Inches(0.9), Inches(1.05), Inches(11.5), Inches(1.4))
tf2 = tb_title2.text_frame
tf2.word_wrap = True
tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

p2_1 = tf2.paragraphs[0]
p2_1.text = "The High-ROI Solution: Scalable Risk Triage & Monetization"
p2_1.font.size = Pt(30)
p2_1.font.bold = True
p2_1.font.color.rgb = COLOR_DARK_FOREST
p2_1.space_after = Pt(4)

p2_2 = tf2.add_paragraph()
p2_2.text = "Enterprise Architecture, Direct Municipal ROI, and Nationwide SaaS Expansion"
p2_2.font.size = Pt(16)
p2_2.font.color.rgb = RGBColor(0x56, 0x64, 0x5C)

# 3 Cards on Slide 2
make_card(
    slide2, Inches(0.9), Inches(2.55), Inches(3.65), Inches(3.8),
    "Enterprise Cloud Moat", RGBColor(0xE5, 0xF0, 0xE1), COLOR_SAGE_ACCENT,
    "☁ Cloud-Native Stack",
    [
        "Google Cloud Storage: Automated public blob bucket (tree-hazard-images-503718)",
        "Google Maps Platform: Dual-layer geocoding & static street inspections",
        "Supabase PostgreSQL: Immutable municipal audit trail & arborist logs",
        "Zero-delay Node.js + React 19: Tested, 100% CI pass rate, production-ready"
    ],
    is_dark=False
)

make_card(
    slide2, Inches(4.84), Inches(2.55), Inches(3.65), Inches(3.8),
    "Direct Financial ROI", RGBColor(0xFD, 0xF0, 0xE2), COLOR_TERRACOTTA,
    "💰 Massive Cost Savings",
    [
        "$2.4M saved per 100k residents in avoided storm grid repair & liability",
        "60% reduction in truck rolls: Arborists dispatched only with verified photo evidence",
        "Proximity route clustering: Slashes crew travel fuel and deadhead miles",
        "Fair & equitable: Transparent math protects underserved neighborhoods"
    ],
    is_dark=False
)

make_card(
    slide2, Inches(8.78), Inches(2.55), Inches(3.65), Inches(3.8),
    "Market Expansion", RGBColor(0xED, 0xE7, 0xDF), RGBColor(0x6D, 0x56, 0x48),
    "🚀 $1.8B Municipal TAM",
    [
        "4,200+ North American cities operate on the same broken 311 tree queue",
        "Plug-and-play ArcGIS connector allows turnkey municipal onboarding in < 48 hours",
        "Insurance & utility partnership upside: powerline vegetation monitoring contracts",
        "Proven & validated today in Halifax Regional Municipality"
    ],
    is_dark=False
)

# Thank You Strip on Slide 2
strip = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(6.52), Inches(11.53), Inches(0.68))
strip.fill.solid()
strip.fill.fore_color.rgb = COLOR_DARK_FOREST
strip.line.color.rgb = COLOR_SAGE_ACCENT

tf_strip = strip.text_frame
tf_strip.word_wrap = True
p_st1 = tf_strip.paragraphs[0]
p_st1.text = "🌲 Thank You! Let's Protect Cities & Power Grids Together.  |  Questions & Demonstration • Team Lambda Legends"
p_st1.font.size = Pt(12)
p_st1.font.bold = True
p_st1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
p_st1.alignment = PP_ALIGN.CENTER

# Save presentation
output_path = os.path.join(os.getcwd(), "CanopyGuard_Executive_Pitch.pptx")
prs.save(output_path)
print(f"Successfully generated: {output_path}")
